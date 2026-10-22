import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def generate_all_orders_excel(orders_qs):
    """
    Generates a beautifully formatted, multi-sheet Excel (.xlsx) workbook
    containing all store orders with both order summaries and itemized product breakdown.
    """
    wb = openpyxl.Workbook()

    # Style constants
    header_fill = PatternFill(start_color="9E2C18", end_color="9E2C18", fill_type="solid")
    sub_header_fill = PatternFill(start_color="3B160E", end_color="3B160E", fill_type="solid")
    accent_fill = PatternFill(start_color="FFF3E0", end_color="FFF3E0", fill_type="solid")
    total_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
    confirmed_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    bold_font = Font(name="Calibri", size=11, bold=True, color="2A160D")
    title_font = Font(name="Calibri", size=14, bold=True, color="9E2C18")
    regular_font = Font(name="Calibri", size=10, color="2A160D")
    confirmed_font = Font(name="Calibri", size=10, bold=True, color="15803D")

    thin_border_side = Side(border_style="thin", color="E5D5C5")
    cell_border = Border(
        left=thin_border_side,
        right=thin_border_side,
        top=thin_border_side,
        bottom=thin_border_side,
    )
    total_border = Border(
        top=Side(border_style="thin", color="9E2C18"),
        bottom=Side(border_style="double", color="9E2C18"),
    )

    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    align_right = Alignment(horizontal="right", vertical="center")

    # ==========================================
    # SHEET 1: All Orders Summary
    # ==========================================
    ws_orders = wb.active
    ws_orders.title = "All Orders"
    ws_orders.views.sheetView[0].showGridLines = True

    # Title block
    ws_orders.merge_cells("A1:M1")
    title_cell = ws_orders["A1"]
    title_cell.value = "THASWIKA CRACKERS SIVAKASI - MASTER ORDERS REPORT"
    title_cell.font = title_font
    title_cell.alignment = Alignment(horizontal="left", vertical="center")
    ws_orders.row_dimensions[1].height = 28

    headers = [
        "S.No",
        "Order #",
        "Order Date",
        "Customer Username",
        "Customer Name",
        "Phone Number",
        "Delivery Address",
        "Total Items Qty",
        "Products Breakdown",
        "Total Amount (₹)",
        "Order Status",
        "Admin Confirmed",
        "Customer Notes",
    ]

    ws_orders.append([])  # Blank row 2
    ws_orders.append(headers)  # Row 3
    ws_orders.row_dimensions[3].height = 26

    for col_idx, _ in enumerate(headers, start=1):
        c = ws_orders.cell(row=3, column=col_idx)
        c.fill = header_fill
        c.font = header_font
        c.alignment = align_center
        c.border = cell_border

    orders_list = list(orders_qs)
    total_sales_value = 0
    total_products_count = 0

    current_row = 4
    for idx, order in enumerate(orders_list, start=1):
        items = list(order.items.all())
        items_summary = "; ".join(f"{i.product_name} x {i.quantity}" for i in items)
        order_qty = sum(i.quantity for i in items)
        total_products_count += order_qty
        total_sales_value += float(order.total or 0)

        created_str = order.created_at.strftime("%Y-%m-%d %H:%M") if order.created_at else ""
        username_str = order.user.username if order.user else "Guest"
        is_confirmed = "YES" if (order.status or "").lower() == "confirmed" else ("COMPLETED" if (order.status or "").lower() == "completed" else "NO")

        row_values = [
            idx,
            order.order_number,
            created_str,
            username_str,
            order.customer_name,
            str(order.phone),
            order.address,
            order_qty,
            items_summary,
            float(order.total or 0),
            (order.status or "new").upper(),
            is_confirmed,
            order.notes or "",
        ]
        ws_orders.append(row_values)
        ws_orders.row_dimensions[current_row].height = 22

        # Style data row
        for col_idx in range(1, len(row_values) + 1):
            cell = ws_orders.cell(row=current_row, column=col_idx)
            cell.font = regular_font
            cell.border = cell_border

            # Column specific alignments & formats
            if col_idx in (1, 8):
                cell.alignment = align_center
            elif col_idx in (2, 3, 4, 6):
                cell.alignment = align_center
            elif col_idx == 10:
                cell.alignment = align_right
                cell.number_format = "₹#,##0.00"
            elif col_idx == 11:
                cell.alignment = align_center
                if (order.status or "").lower() == "confirmed":
                    cell.fill = confirmed_fill
                    cell.font = confirmed_font
            elif col_idx == 12:
                cell.alignment = align_center
                if is_confirmed == "YES":
                    cell.fill = confirmed_fill
                    cell.font = confirmed_font
            else:
                cell.alignment = align_left

        current_row += 1

    # Totals Row
    ws_orders.append([
        "TOTAL",
        f"{len(orders_list)} Orders",
        "",
        "",
        "",
        "",
        "",
        total_products_count,
        "",
        total_sales_value,
        "",
        "",
        "",
    ])
    ws_orders.row_dimensions[current_row].height = 24
    for col_idx in range(1, 14):
        c = ws_orders.cell(row=current_row, column=col_idx)
        c.font = bold_font
        c.fill = total_fill
        c.border = total_border
        if col_idx in (1, 2, 8):
            c.alignment = align_center
        elif col_idx == 10:
            c.alignment = align_right
            c.number_format = "₹#,##0.00"

    # ==========================================
    # SHEET 2: Itemized Products Breakdown
    # ==========================================
    ws_items = wb.create_sheet(title="Itemized Products")
    ws_items.views.sheetView[0].showGridLines = True

    ws_items.merge_cells("A1:I1")
    items_title = ws_items["A1"]
    items_title.value = "ALL ORDER ITEMS DETAILED BREAKDOWN"
    items_title.font = title_font
    items_title.alignment = Alignment(horizontal="left", vertical="center")
    ws_items.row_dimensions[1].height = 26

    item_headers = [
        "Order #",
        "Date",
        "Customer Name",
        "Phone",
        "Product Code",
        "Product Description",
        "Quantity",
        "Unit Price (₹)",
        "Subtotal (₹)",
    ]
    ws_items.append([])
    ws_items.append(item_headers)
    ws_items.row_dimensions[3].height = 24

    for col_idx, _ in enumerate(item_headers, start=1):
        c = ws_items.cell(row=3, column=col_idx)
        c.fill = sub_header_fill
        c.font = header_font
        c.alignment = align_center
        c.border = cell_border

    item_row_idx = 4
    for order in orders_list:
        created_str = order.created_at.strftime("%Y-%m-%d %H:%M") if order.created_at else ""
        for item in order.items.all():
            row_data = [
                order.order_number,
                created_str,
                order.customer_name,
                str(order.phone),
                str(item.product_code or "-"),
                item.product_name,
                item.quantity,
                float(item.unit_price or 0),
                float(item.subtotal or 0),
            ]
            ws_items.append(row_data)
            ws_items.row_dimensions[item_row_idx].height = 20

            for col_idx in range(1, len(row_data) + 1):
                cell = ws_items.cell(row=item_row_idx, column=col_idx)
                cell.font = regular_font
                cell.border = cell_border
                if col_idx in (1, 2, 4, 5, 7):
                    cell.alignment = align_center
                elif col_idx in (8, 9):
                    cell.alignment = align_right
                    cell.number_format = "₹#,##0.00"
                else:
                    cell.alignment = align_left

            item_row_idx += 1

    # Adjust column widths automatically for both sheets
    for ws in [ws_orders, ws_items]:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            # Skip title row in width calculation
            for cell in list(col)[2:]:
                val_str = str(cell.value or "")
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 11)

    # Specific column width overrides for better readability
    ws_orders.column_dimensions["A"].width = 7
    ws_orders.column_dimensions["B"].width = 22
    ws_orders.column_dimensions["C"].width = 18
    ws_orders.column_dimensions["D"].width = 18
    ws_orders.column_dimensions["E"].width = 22
    ws_orders.column_dimensions["F"].width = 16
    ws_orders.column_dimensions["G"].width = 30
    ws_orders.column_dimensions["H"].width = 15
    ws_orders.column_dimensions["I"].width = 40
    ws_orders.column_dimensions["J"].width = 18
    ws_orders.column_dimensions["K"].width = 16
    ws_orders.column_dimensions["L"].width = 16
    ws_orders.column_dimensions["M"].width = 25

    ws_items.column_dimensions["A"].width = 22
    ws_items.column_dimensions["B"].width = 18
    ws_items.column_dimensions["C"].width = 22
    ws_items.column_dimensions["D"].width = 16
    ws_items.column_dimensions["E"].width = 14
    ws_items.column_dimensions["F"].width = 35
    ws_items.column_dimensions["G"].width = 12
    ws_items.column_dimensions["H"].width = 16
    ws_items.column_dimensions["I"].width = 16

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
